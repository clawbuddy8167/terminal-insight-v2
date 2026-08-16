"""Stage 5 — Excel: write parameter workbook from resolutions."""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from rich.console import Console

from terminal_insight_v2.category.loader import load_category_schema
from terminal_insight_v2.config import settings
from terminal_insight_v2.models.stage import StageContext, StageResult
from terminal_insight_v2.stages.base import Stage
from terminal_insight_v2.store.db import connection, utc_now

console = Console()


def _load_resolutions(run_id: str) -> List[Dict[str, Any]]:
    sql = """
      SELECT r.device_id, r.field_name, r.display_text, r.policy,
             r.selected_assertion_ids, r.conflict, r.locked,
             d.brand, d.model, d.release_date
      FROM resolutions r
      JOIN devices d ON d.device_id = r.device_id
      WHERE r.run_id = ?
      ORDER BY d.brand, d.model, r.field_name
    """
    with connection(readonly=True) as conn:
        return [dict(row) for row in conn.execute(sql, (run_id,)).fetchall()]


def _load_assertion_urls(run_id: str) -> Dict[str, str]:
    sql = """
      SELECT a.assertion_id, s.url
      FROM field_assertions a
      JOIN sources s ON s.source_id = a.source_id
      WHERE a.run_id = ?
    """
    with connection(readonly=True) as conn:
        return {r["assertion_id"]: r["url"] for r in conn.execute(sql, (run_id,)).fetchall()}


class ExcelStage(Stage):
    name = "excel"

    def prereq(self, ctx: StageContext) -> None:
        resolved = Path(ctx.run_dir) / "resolved.json"
        rows = _load_resolutions(ctx.run_id)
        if not rows and not resolved.exists():
            raise FileNotFoundError(
                f"缺少 resolutions（run={ctx.run_id}）；请先跑 verify"
            )

    def run(self, ctx: StageContext) -> StageResult:
        schema = load_category_schema(ctx.category, persist=False)
        rows = _load_resolutions(ctx.run_id)
        url_by_aid = _load_assertion_urls(ctx.run_id)

        # device_id -> field -> resolution
        by_device: Dict[str, Dict[str, Dict[str, Any]]] = defaultdict(dict)
        device_meta: Dict[str, Dict[str, Any]] = {}
        for r in rows:
            by_device[r["device_id"]][r["field_name"]] = r
            device_meta[r["device_id"]] = {
                "brand": r["brand"],
                "model": r["model"],
                "release_date": r.get("release_date") or "",
            }

        columns = sorted(schema.output_columns, key=lambda c: c.order)
        if not columns:
            columns = [
                type("C", (), {"field_name": f.name, "header": f.display_name or f.name, "width": 24, "wrap": True})()
                for f in schema.fields
            ]

        wb = Workbook()
        # Sheet 1: parameters
        ws = wb.active
        ws.title = f"{schema.label_zh}参数"[:31]
        headers = [c.header or c.field_name for c in columns]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        gold_mode = settings.report_mode == "gold"
        for did, fields in by_device.items():
            row_vals: List[str] = []
            for col in columns:
                fr = fields.get(col.field_name)
                if not fr:
                    # identity fallback
                    meta = device_meta[did]
                    if col.field_name == "品牌":
                        row_vals.append(meta["brand"])
                    elif col.field_name == "机型":
                        row_vals.append(meta["model"])
                    elif col.field_name == "发布时间":
                        row_vals.append(meta["release_date"])
                    else:
                        row_vals.append("")
                    continue
                if gold_mode and fr.get("policy") not in ("locked", "unanimous", "prefer_official"):
                    # gold mode: only locked / high-confidence; leave multi_claim empty for now
                    if fr.get("policy") == "multi_claim" or fr.get("policy") == "missing":
                        row_vals.append("")
                        continue
                row_vals.append(fr.get("display_text") or "")
            ws.append(row_vals)

        for idx, col in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = col.width or 22
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Sheet 2: field sources
        ws2 = wb.create_sheet("字段信源")
        ws2.append(["品牌", "机型", "字段", "展示值", "策略", "冲突", "信源URL"])
        for cell in ws2[1]:
            cell.font = Font(bold=True)
        for did, fields in by_device.items():
            meta = device_meta[did]
            for fname, fr in fields.items():
                urls: List[str] = []
                try:
                    aids = json.loads(fr.get("selected_assertion_ids") or "[]")
                except Exception:
                    aids = []
                for aid in aids:
                    u = url_by_aid.get(aid)
                    if u:
                        urls.append(u)
                ws2.append(
                    [
                        meta["brand"],
                        meta["model"],
                        fname,
                        fr.get("display_text") or "",
                        fr.get("policy") or "",
                        "是" if fr.get("conflict") else "否",
                        " | ".join(urls),
                    ]
                )

        # Sheet 3: conflicts / issues from review_queue if present
        ws3 = wb.create_sheet("冲突与校验")
        ws3.append(["严重度", "品牌", "机型", "字段", "问题", "说明", "URL"])
        for cell in ws3[1]:
            cell.font = Font(bold=True)
        review_path = Path(ctx.run_dir) / "review_queue.json"
        if review_path.exists():
            data = json.loads(review_path.read_text(encoding="utf-8"))
            for issue in data.get("issues") or []:
                ws3.append(
                    [
                        issue.get("severity", ""),
                        issue.get("brand", ""),
                        issue.get("model", ""),
                        issue.get("field_name", ""),
                        issue.get("issue", ""),
                        issue.get("notes", ""),
                        issue.get("source_url", ""),
                    ]
                )
            for q in data.get("queue") or []:
                if q.get("reason") == "conflict":
                    ws3.append(
                        [
                            "warning",
                            q.get("brand", ""),
                            q.get("model", ""),
                            q.get("field_name", ""),
                            "conflict",
                            q.get("display_text", ""),
                            "",
                        ]
                    )

        # Vote sheet placeholder; Vote stage fills rows
        ws4 = wb.create_sheet("亮点投票")
        ws4.append(["参考系", "取景框", "热点", "品牌", "机型", "说明"])

        run_dir = Path(ctx.run_dir)
        out_name = f"{ctx.category}_{ctx.date_start.replace('-', '')}-{ctx.date_end.replace('-', '')}.xlsx"
        if not ctx.date_start:
            out_name = f"{ctx.category}_{ctx.run_id}.xlsx"
        out_path = run_dir / out_name
        wb.save(out_path)
        console.log(f"[excel] wrote {out_path.name} rows={len(by_device)}")

        # Also dump a compact resolved table json
        table_path = run_dir / "excel_table.json"
        table_path.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "run_id": ctx.run_id,
                    "category": ctx.category,
                    "generated_at": utc_now(),
                    "columns": [c.header or c.field_name for c in columns],
                    "row_count": len(by_device),
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

        return StageResult(
            stage=self.name,
            status="ok",
            output_revision=f"excel:{ctx.run_id}:n{len(by_device)}",
            metrics={"rows": len(by_device), "columns": len(columns)},
            artifacts={
                "excel": str(out_path),
                "excel_table.json": str(table_path),
            },
        )
