"""Excel export for settlement artifacts (research layer, not facts SoT)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .models import SettlementArtifact

SHEET_NAME = "假说结算"
HEADERS = [
    "假说ID",
    "命题",
    "判决",
    "说明",
    "各季支持/削弱",
    "代表热点",
]
_WIDTHS = (10, 40, 12, 28, 36, 40)

OEM_SHEET = "OEM切片"
OEM_HEADERS = ["run_id", "热点", "品牌", "机型", "相关度", "行动启示", "为什么关注", "边界"]
_OEM_WIDTHS = (22, 24, 12, 16, 8, 32, 28, 24)


def write_settlement_excel(path: Path, artifact: SettlementArtifact) -> Path:
    wb = Workbook()
    default = wb.active
    if default is not None and default.title == "Sheet":
        wb.remove(default)

    ws = wb.create_sheet(SHEET_NAME, 0)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in artifact.settlements:
        quarter_bits = []
        for q in row.by_quarter:
            c = q.counts
            quarter_bits.append(f"{q.run_id}: +{c.support}/-{c.weaken}")
        reps = "；".join(
            f"[{r.run_id}] {r.title} ({r.stance})" for r in row.representatives[:3]
        )
        ws.append(
            [
                row.hypothesis_id,
                row.statement,
                row.verdict,
                row.note,
                " | ".join(quarter_bits),
                reps,
            ]
        )
    for idx, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ows = wb.create_sheet(OEM_SHEET)
    ows.append(OEM_HEADERS)
    for cell in ows[1]:
        cell.font = Font(bold=True)
    for item in artifact.oem_slice:
        ows.append(
            [
                item.run_id,
                item.title,
                item.brand,
                item.model,
                item.relevance,
                item.action_implication,
                item.why_it_matters,
                item.caveat,
            ]
        )
    for idx, width in enumerate(_OEM_WIDTHS, start=1):
        ows.column_dimensions[get_column_letter(idx)].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
