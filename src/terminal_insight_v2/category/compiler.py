"""Compile category Excel workbooks into CategorySchema."""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import load_workbook

from terminal_insight_v2.identity import content_sha256
from terminal_insight_v2.store.db import utc_now

from .schema import (
    CategorySchema,
    DiscoveryConfig,
    FieldDef,
    OutputColumn,
    SourceDef,
)
from .validators_registry import KNOWN_VALIDATORS

_CATEGORY_ID_RE = re.compile(r"^[a-z][a-z0-9_]*$")
_REQUIRED_SHEETS = ("品类配置", "字段定义", "输出模板")
_FIELD_COLS = {
    "field_name",
    "display_name",
    "description",
    "data_type",
    "required",
    "is_identity",
    "key_spec",
    "example",
    "search_keywords",
    "validators",
    "conflict_policy",
    "empty_display",
    "notes",
}
_OUTPUT_COLS = {"column_order", "field_name", "header", "width", "wrap"}
_SOURCE_COLS = {
    "source_key",
    "name_zh",
    "tier",
    "source_kind",
    "domains",
    "priority",
    "auto_search",
    "enabled",
}


class CategoryCompileError(Exception):
    def __init__(self, errors: Sequence[str], *, path: str = "") -> None:
        self.errors = list(errors)
        self.path = path
        head = path or "template"
        body = "\n".join(f"  - {e}" for e in self.errors)
        super().__init__(f"[category-compile] {head} 失败：\n{body}")


def _cell(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _as_bool(v: Any, *, default: bool = False) -> bool:
    s = _cell(v).lower()
    if not s:
        return default
    if s in {"true", "1", "yes", "y", "是"}:
        return True
    if s in {"false", "0", "no", "n", "否"}:
        return False
    raise ValueError(f"无法解析布尔值: {v!r}")


def _split_list(v: Any) -> List[str]:
    s = _cell(v)
    if not s:
        return []
    parts = re.split(r"[;；|]", s)
    return [p.strip() for p in parts if p.strip()]


def _sheet_map(wb) -> Dict[str, Any]:
    return {str(name).strip(): wb[name] for name in wb.sheetnames}


def _read_kv_sheet(ws) -> Dict[str, str]:
    out: Dict[str, str] = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        key = _cell(row[0])
        val = _cell(row[1]) if len(row) > 1 else ""
        if key:
            out[key] = val
    return out


def _header_rows(ws) -> Tuple[List[str], List[Tuple[int, Dict[str, Any]]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [_cell(h).lower() for h in rows[0]]
    # keep original case mapping via lower
    data: List[Tuple[int, Dict[str, Any]]] = []
    for i, row in enumerate(rows[1:], start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        item: Dict[str, Any] = {}
        for idx, h in enumerate(headers):
            if not h:
                continue
            item[h] = row[idx] if idx < len(row) else None
        data.append((i, item))
    return headers, data


def compile_template(
    path: str | Path,
    *,
    persist: bool = False,
    memory_root: Optional[Path] = None,
) -> CategorySchema:
    path = Path(path).expanduser().resolve()
    errors: List[str] = []
    if not path.is_file():
        raise CategoryCompileError([f"文件不存在: {path}"], path=str(path))

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = _sheet_map(wb)
        for name in _REQUIRED_SHEETS:
            if name not in sheets:
                errors.append(f"缺少必填 Sheet「{name}」")
        if errors:
            raise CategoryCompileError(errors, path=path.name)

        cfg = _read_kv_sheet(sheets["品类配置"])
        category_id = cfg.get("category_id", "")
        if not category_id:
            errors.append("Sheet「品类配置」：缺少 category_id")
        elif not _CATEGORY_ID_RE.match(category_id):
            errors.append(
                f"Sheet「品类配置」：category_id={category_id!r} 须匹配 [a-z][a-z0-9_]*"
            )
        elif category_id != path.stem:
            errors.append(
                f"Sheet「品类配置」：category_id={category_id} 与文件名 {path.stem} 不一致"
            )

        for key in ("label_zh", "template_version", "discovery_product"):
            if not cfg.get(key):
                errors.append(f"Sheet「品类配置」：缺少必填键 {key}")

        # fields
        headers, field_rows = _header_rows(sheets["字段定义"])
        unknown = set(headers) - _FIELD_COLS - {""}
        if unknown:
            errors.append(
                f"Sheet「字段定义」：未知列 {sorted(unknown)}"
            )
        for col in ("field_name", "description", "data_type", "required", "is_identity", "example"):
            if col not in headers:
                errors.append(f"Sheet「字段定义」：缺少必填列 {col}")

        fields: List[FieldDef] = []
        seen_names: set[str] = set()
        for row_no, item in field_rows:
            name = _cell(item.get("field_name"))
            desc = _cell(item.get("description"))
            example = _cell(item.get("example"))
            if not name:
                errors.append(f"Sheet「字段定义」第 {row_no} 行：缺少 field_name")
                continue
            if name in seen_names:
                errors.append(
                    f"Sheet「字段定义」第 {row_no} 行：field_name={name} 重复"
                )
            seen_names.add(name)
            if not desc:
                errors.append(
                    f"Sheet「字段定义」第 {row_no} 行：缺少必填列值 description"
                    f"（field_name={name}）"
                )
            if not example:
                errors.append(
                    f"Sheet「字段定义」第 {row_no} 行：缺少必填列值 example"
                    f"（field_name={name}）"
                )
            try:
                required = _as_bool(item.get("required"), default=False)
                is_identity = _as_bool(item.get("is_identity"), default=False)
                key_spec = _as_bool(item.get("key_spec"), default=False)
            except ValueError as exc:
                errors.append(f"Sheet「字段定义」第 {row_no} 行：{exc}")
                continue
            dtype = _cell(item.get("data_type")) or "text"
            policy = _cell(item.get("conflict_policy")) or "multi_claim"
            validators = _split_list(item.get("validators"))
            bad_v = [v for v in validators if v not in KNOWN_VALIDATORS]
            if bad_v:
                errors.append(
                    f"Sheet「字段定义」第 {row_no} 行：未知 validators {bad_v}"
                )
            if dtype not in {"string", "text", "price", "date", "enum", "boolean"}:
                errors.append(
                    f"Sheet「字段定义」第 {row_no} 行：非法 data_type={dtype}"
                )
            if policy not in {"multi_claim", "prefer_official", "prefer_latest"}:
                errors.append(
                    f"Sheet「字段定义」第 {row_no} 行：非法 conflict_policy={policy}"
                )
            if not desc or not example:
                continue
            fields.append(
                FieldDef(
                    name=name,
                    display_name=_cell(item.get("display_name")) or name,
                    description=desc,
                    data_type=dtype,  # type: ignore[arg-type]
                    required=required,
                    is_identity=is_identity,
                    key_spec=key_spec,
                    example=example,
                    search_keywords=_split_list(item.get("search_keywords")),
                    validators=validators,
                    conflict_policy=policy,  # type: ignore[arg-type]
                    empty_display=_cell(item.get("empty_display")),
                    notes=_cell(item.get("notes")),
                )
            )

        if "品牌" not in seen_names or "机型" not in seen_names:
            errors.append("Sheet「字段定义」：必须包含身份字段 品牌 与 机型")

        # output
        o_headers, out_rows = _header_rows(sheets["输出模板"])
        unknown_o = set(o_headers) - _OUTPUT_COLS - {""}
        if unknown_o:
            errors.append(f"Sheet「输出模板」：未知列 {sorted(unknown_o)}")
        for col in ("column_order", "field_name"):
            if col not in o_headers:
                errors.append(f"Sheet「输出模板」：缺少必填列 {col}")

        output_columns: List[OutputColumn] = []
        orders: set[int] = set()
        for row_no, item in out_rows:
            fname = _cell(item.get("field_name"))
            if not fname:
                errors.append(f"Sheet「输出模板」第 {row_no} 行：缺少 field_name")
                continue
            if fname not in seen_names:
                errors.append(
                    f"Sheet「输出模板」：field_name={fname} 未在「字段定义」中声明"
                )
            try:
                order = int(float(_cell(item.get("column_order"))))
            except (TypeError, ValueError):
                errors.append(
                    f"Sheet「输出模板」第 {row_no} 行：column_order 非法"
                )
                continue
            if order in orders:
                errors.append(
                    f"Sheet「输出模板」第 {row_no} 行：column_order={order} 重复"
                )
            orders.add(order)
            width_raw = _cell(item.get("width"))
            width = int(float(width_raw)) if width_raw else None
            try:
                wrap = _as_bool(item.get("wrap"), default=True)
            except ValueError as exc:
                errors.append(f"Sheet「输出模板」第 {row_no} 行：{exc}")
                wrap = True
            output_columns.append(
                OutputColumn(
                    order=order,
                    field_name=fname,
                    header=_cell(item.get("header")) or fname,
                    width=width,
                    wrap=wrap,
                )
            )
        output_columns.sort(key=lambda c: c.order)

        # optional sources
        sources: List[SourceDef] = []
        if "信源配置" in sheets:
            s_headers, s_rows = _header_rows(sheets["信源配置"])
            unknown_s = set(s_headers) - _SOURCE_COLS - {""}
            if unknown_s:
                errors.append(f"Sheet「信源配置」：未知列 {sorted(unknown_s)}")
            for row_no, item in s_rows:
                key = _cell(item.get("source_key"))
                if not key:
                    continue
                tier = _cell(item.get("tier"))
                if tier not in {"P0", "P1", "P2", "P3"}:
                    errors.append(
                        f"Sheet「信源配置」第 {row_no} 行：非法 tier={tier}"
                    )
                    continue
                try:
                    priority = int(float(_cell(item.get("priority")) or "100"))
                    auto_search = _as_bool(item.get("auto_search"), default=True)
                    enabled = _as_bool(item.get("enabled"), default=True)
                except ValueError as exc:
                    errors.append(f"Sheet「信源配置」第 {row_no} 行：{exc}")
                    continue
                sources.append(
                    SourceDef(
                        source_key=key,
                        name_zh=_cell(item.get("name_zh")) or key,
                        tier=tier,  # type: ignore[arg-type]
                        source_kind=_cell(item.get("source_kind")) or "named",
                        domains=_split_list(item.get("domains")),
                        priority=priority,
                        auto_search=auto_search,
                        enabled=enabled,
                    )
                )

        if errors:
            raise CategoryCompileError(errors, path=path.name)

        identity_fields = _split_list(cfg.get("identity_fields")) or [
            "品牌",
            "机型",
            "发布时间",
        ]
        try:
            target_sources = int(float(cfg.get("target_source_count") or "10"))
            spec_req = _as_bool(
                cfg.get("official_spec_required_if_exists"), default=True
            )
            mall_req = _as_bool(
                cfg.get("official_mall_required_if_exists"), default=True
            )
        except ValueError as exc:
            raise CategoryCompileError(
                [f"Sheet「品类配置」：{exc}"], path=path.name
            ) from exc

        schema = CategorySchema(
            category_id=category_id,
            label_zh=cfg["label_zh"],
            label_en=cfg.get("label_en", ""),
            template_version=cfg["template_version"],
            region_default=cfg.get("region_default") or "CN",
            identity_fields=identity_fields,
            official_spec_required_if_exists=spec_req,
            official_mall_required_if_exists=mall_req,
            target_source_count=target_sources,
            discovery=DiscoveryConfig(
                product=cfg["discovery_product"],
                extra=cfg.get("discovery_extra", ""),
                prompt_extra=cfg.get("discovery_prompt_extra", ""),
                search_brands=_split_list(cfg.get("search_brands")),
            ),
            fields=fields,
            output_columns=output_columns,
            sources=sources,
            compiled_at=utc_now(),
            template_path=str(path),
        )
        canonical = json.dumps(
            schema.to_canonical_dict(),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
        schema.fingerprint = content_sha256(canonical)

        if persist:
            from terminal_insight_v2.store.db import ensure_memory_dirs, transaction

            root = memory_root
            ensure_memory_dirs(root=root)
            out_dir = (root or Path(".")) 
            # resolve via settings if needed
            from terminal_insight_v2.config import settings

            base = root or settings.memory_root()
            compiled_dir = base / "compiled"
            compiled_dir.mkdir(parents=True, exist_ok=True)
            out_path = compiled_dir / f"{schema.category_id}@{schema.fingerprint[:16]}.json"
            out_path.write_text(schema.dumps(), encoding="utf-8")
            with transaction(root=root) as conn:
                conn.execute(
                    """
                    INSERT INTO category_schemas(
                      schema_id, category, version, fingerprint,
                      template_path, compiled_json, compiled_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(fingerprint) DO UPDATE SET
                      compiled_json=excluded.compiled_json,
                      compiled_at=excluded.compiled_at,
                      template_path=excluded.template_path
                    """,
                    (
                        f"schema:{schema.category_id}:{schema.fingerprint[:16]}",
                        schema.category_id,
                        schema.template_version,
                        schema.fingerprint,
                        str(path),
                        schema.dumps(),
                        schema.compiled_at,
                    ),
                )
        return schema
    finally:
        wb.close()


__all__ = ["CategoryCompileError", "compile_template"]
