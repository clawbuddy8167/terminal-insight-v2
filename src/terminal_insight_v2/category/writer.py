"""Write category Excel templates from structured dicts (bootstrap / scaffold)."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Sequence

from openpyxl import Workbook


def write_category_workbook(path: str | Path, spec: Dict[str, Any]) -> Path:
    """Write a multi-sheet category template.

    ``spec`` keys: category (kv dict), fields (list[dict]), output (list[dict]),
    optional sources (list[dict]).
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "品类配置"
    for i, (k, v) in enumerate((spec.get("category") or {}).items(), start=1):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)

    field_headers = [
        "field_name",
        "display_name",
        "description",
        "data_type",
        "required",
        "is_identity",
        "key_spec",
        "example",
        "search_keywords",
        "validators",
        "conflict_policy",
        "empty_display",
        "notes",
    ]
    ws_f = wb.create_sheet("字段定义")
    for c, h in enumerate(field_headers, start=1):
        ws_f.cell(1, c, h)
    for r, row in enumerate(spec.get("fields") or [], start=2):
        for c, h in enumerate(field_headers, start=1):
            ws_f.cell(r, c, row.get(h, ""))

    out_headers = ["column_order", "field_name", "header", "width", "wrap"]
    ws_o = wb.create_sheet("输出模板")
    for c, h in enumerate(out_headers, start=1):
        ws_o.cell(1, c, h)
    for r, row in enumerate(spec.get("output") or [], start=2):
        for c, h in enumerate(out_headers, start=1):
            ws_o.cell(r, c, row.get(h, ""))

    sources: Sequence[Dict[str, Any]] = spec.get("sources") or []
    if sources:
        src_headers = [
            "source_key",
            "name_zh",
            "tier",
            "source_kind",
            "domains",
            "priority",
            "auto_search",
            "enabled",
        ]
        ws_s = wb.create_sheet("信源配置")
        for c, h in enumerate(src_headers, start=1):
            ws_s.cell(1, c, h)
        for r, row in enumerate(sources, start=2):
            for c, h in enumerate(src_headers, start=1):
                ws_s.cell(r, c, row.get(h, ""))

    wb.save(path)
    return path


def fields_to_output(fields: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    out = []
    for i, f in enumerate(fields, start=1):
        out.append(
            {
                "column_order": i,
                "field_name": f["field_name"],
                "header": f.get("display_name") or f["field_name"],
                "width": 24 if f.get("is_identity") in (True, "true") else 40,
                "wrap": "true",
            }
        )
    return out


__all__ = ["fields_to_output", "write_category_workbook"]
