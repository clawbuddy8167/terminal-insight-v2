"""Convert the 2026 collected-fields workbook into per-category resolved.json.

Does not search, extract, or verify. Maps 手机/眼镜/手表 sheets onto the
pipeline category field names so Vote and quadrant can run from the table.
"""
from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from openpyxl import load_workbook

from terminal_insight_v2.category import load_category_schema
from terminal_insight_v2.store.db import utc_now
from terminal_insight_v2.store.ledger import begin_run

HEADER_ROW = 4
DATA_START_ROW = 5
JOIN = "；"

SHEETS: Dict[str, str] = {
    "手机": "phone",
    "眼镜": "glasses",
    "手表": "watch",
}

IDENTITY = frozenset({"品牌", "机型", "发布时间"})
SKIP_DATA_STATUS = frozenset({"candidate_only"})

# pipeline field → source columns (left-to-right join).
PHONE_MAP: Dict[str, Tuple[str, ...]] = {
    "品牌": ("品牌",),
    "机型": ("产品名称",),
    "发布时间": ("中国区发布日期",),
    "起售价": ("中国区起售价",),
    "电池容量": ("电池容量",),
    "充电功率": ("有线充电", "无线充电"),
    "防水和防护": ("防尘防水等级",),
    "核心算力与存储": ("芯片", "内存规格", "存储规格"),
    "前沿通信能力": ("蜂窝网络", "Wi-Fi版本", "蓝牙版本", "NFC支持"),
    "屏幕显示技术": ("主屏尺寸", "屏幕类型", "屏幕分辨率", "最高刷新率"),
    "智能传感与交互": ("指纹识别技术", "指纹安装位置"),
    "AI底层能力": ("AI功能", "AI功能标签", "AI最高能力阶段"),
    "核心宣发亮点与差异化功能": (
        "摄像头摘要",
        "卫星通信（是否支持）",
        "主要竞争路线",
        "产品形态",
    ),
}

WATCH_MAP: Dict[str, Tuple[str, ...]] = {
    "品牌": ("品牌",),
    "机型": ("产品名称",),
    "发布时间": ("中国区发布日期",),
    "起售价": ("中国区起售价",),
    "表壳尺寸与材质": ("高度", "宽度", "厚度", "重量", "产品形态"),
    "屏幕显示": ("主屏尺寸", "屏幕类型", "屏幕分辨率", "最高刷新率"),
    "续航与充电": ("电池容量", "有线充电", "无线充电"),
    "防水等级": ("防尘防水等级",),
    "传感器与健康": ("AI功能",),
    "系统与连接": ("操作系统", "蜂窝网络", "Wi-Fi版本", "蓝牙版本", "NFC支持"),
    "核心宣发亮点与差异化功能": ("AI功能标签", "主要竞争路线", "摄像头摘要"),
}

GLASSES_MAP: Dict[str, Tuple[str, ...]] = {
    "品牌": ("品牌",),
    "机型": ("产品名称",),
    "发布时间": ("中国区发布日期",),
    "起售价": ("中国区起售价",),
    "是否支持拍摄": ("摄像头摘要", "主摄像素"),
    "是否支持显示": ("屏幕类型", "主屏尺寸", "产品形态"),
    "重量": ("重量",),
    "续航": ("电池容量", "有线充电"),
    "接入AI大模型": ("AI功能",),
    "SOC芯片": ("芯片",),
    "AI应用": ("AI功能标签", "AI最高能力阶段"),
}

FIELD_MAPS = {
    "phone": PHONE_MAP,
    "watch": WATCH_MAP,
    "glasses": GLASSES_MAP,
}

_KIND_BY_SOURCE = {
    "中国区起售价": "price",
    "电池容量": "mah",
    "有线充电": "watt",
    "无线充电": "watt",
    "重量": "weight_g",
}


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def format_cell(value: Any, *, kind: str = "") -> str:
    """Normalize a workbook cell into Vote/quadrant-friendly display_text."""
    if _is_blank(value):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float)):
        if float(value) == int(value):
            n: Any = int(value)
        else:
            n = value
        if kind == "price":
            return f"{n}元"
        if kind == "mah":
            return f"{n}mAh"
        if kind == "watt":
            return f"{n}W"
        if kind == "weight_g":
            return f"{n}g"
        return str(n)
    text = str(value).strip()
    if kind == "price" and re.fullmatch(r"\d+(?:\.\d+)?", text):
        return f"{text}元"
    if kind == "mah" and re.fullmatch(r"\d+(?:\.\d+)?", text):
        return f"{text}mAh"
    if kind == "watt" and re.fullmatch(r"\d+(?:\.\d+)?", text):
        return f"{text}W"
    if kind == "weight_g" and re.fullmatch(r"\d+(?:\.\d+)?", text):
        return f"{text}g"
    return text


def join_sources(
    row: Mapping[str, Any],
    columns: Sequence[str],
    *,
    prefix: bool,
) -> str:
    parts: List[str] = []
    for col in columns:
        kind = _KIND_BY_SOURCE.get(col, "")
        text = format_cell(row.get(col), kind=kind)
        if not text:
            continue
        parts.append(f"{col}：{text}" if prefix and len(columns) > 1 else text)
    return JOIN.join(parts)


def sheet_to_items(
    rows: Iterable[Mapping[str, Any]],
    *,
    category: str,
    include_candidates: bool = False,
) -> List[Dict[str, Any]]:
    """Map tabular rows onto resolved.json items. Skips rows without 品牌+产品名称."""
    fmap = FIELD_MAPS[category]
    items: List[Dict[str, Any]] = []
    for row in rows:
        brand = format_cell(row.get("品牌"))
        model = format_cell(row.get("产品名称"))
        if not brand or not model:
            continue
        status = format_cell(row.get("数据状态"))
        if not include_candidates and status in SKIP_DATA_STATUS:
            continue
        raw_id = format_cell(row.get("产品ID"))
        device_id = raw_id or f"{category}:{brand}:{model}"
        for field_name, columns in fmap.items():
            prefix = field_name not in IDENTITY
            text = join_sources(row, columns, prefix=prefix)
            if not text:
                continue
            items.append(
                {
                    "device_id": device_id,
                    "brand": brand,
                    "model": model,
                    "field_name": field_name,
                    "display_text": text,
                    "policy": "identity" if field_name in IDENTITY else "unanimous",
                    "conflict": False,
                }
            )
    return items


def _header_index(header_row: Sequence[Any]) -> Dict[str, int]:
    out: Dict[str, int] = {}
    for i, cell in enumerate(header_row):
        name = str(cell).strip() if cell is not None else ""
        if name and name not in out:
            out[name] = i
    return out


def read_sheet_rows(path: Path, sheet_name: str) -> List[Dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise FileNotFoundError(f"缺少工作表 {sheet_name!r}: {path}")
    ws = wb[sheet_name]
    header_map: Optional[Dict[str, int]] = None
    rows: List[Dict[str, Any]] = []
    for i, raw in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < HEADER_ROW:
            continue
        if i == HEADER_ROW:
            header_map = _header_index(raw)
            continue
        if not header_map:
            continue
        if not raw or not any(raw):
            continue
        rec: Dict[str, Any] = {}
        for name, idx in header_map.items():
            rec[name] = raw[idx] if idx < len(raw) else None
        rows.append(rec)
    wb.close()
    return rows


def build_resolved(
    *,
    run_id: str,
    category: str,
    items: Sequence[Mapping[str, Any]],
    schema_fingerprint: str = "",
    date_start: str = "",
    date_end: str = "",
    source_xlsx: str = "",
) -> Dict[str, Any]:
    return {
        "schema_version": 1,
        "run_id": run_id,
        "category": category,
        "schema_fingerprint": schema_fingerprint,
        "generated_at": utc_now(),
        "input_revision": f"xlsx:{Path(source_xlsx).name}" if source_xlsx else f"run:{run_id}",
        "date_start": date_start,
        "date_end": date_end,
        "source_xlsx": source_xlsx,
        "items": list(items),
    }


def import_workbook(
    xlsx: Path,
    *,
    date_start: str = "2026-01-01",
    date_end: str = "2026-08-31",
    include_candidates: bool = False,
    sheets: Optional[Mapping[str, str]] = None,
) -> List[Dict[str, Any]]:
    """Create one run per sheet and write resolved.json. Returns run summaries."""
    from terminal_insight_v2.config import settings

    xlsx = Path(xlsx)
    if not xlsx.is_file():
        raise FileNotFoundError(f"找不到采集表: {xlsx}")
    mapping = dict(sheets or SHEETS)
    summaries: List[Dict[str, Any]] = []
    for sheet_name, category in mapping.items():
        schema = load_category_schema(category, persist=False)
        rows = read_sheet_rows(xlsx, sheet_name)
        items = sheet_to_items(
            rows, category=category, include_candidates=include_candidates
        )
        devices = sorted({it["device_id"] for it in items})
        meta = begin_run(
            category=category,
            schema_fingerprint=schema.fingerprint,
            date_start=date_start,
            date_end=date_end,
            input_revision=f"xlsx:{xlsx.name}",
        )
        run_id = str(meta["run_id"])
        run_dir = settings.data_path() / meta["rel_path"]
        payload = build_resolved(
            run_id=run_id,
            category=category,
            items=items,
            schema_fingerprint=schema.fingerprint,
            date_start=date_start,
            date_end=date_end,
            source_xlsx=str(xlsx),
        )
        (run_dir / "resolved.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        summaries.append(
            {
                "sheet": sheet_name,
                "category": category,
                "run_id": run_id,
                "run_dir": str(run_dir),
                "devices": len(devices),
                "items": len(items),
            }
        )
    return summaries
