"""Vote stage: mocked LLM ballots → votes.json + Excel sheet."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from terminal_insight_v2.models.stage import StageContext
from terminal_insight_v2.vote.api import run_vote
from terminal_insight_v2.vote.excel_write import (
    AUDIENCE_HEADERS,
    AUDIENCE_SHEET_NAME,
    HEADERS,
    SHEET_NAME,
)
from terminal_insight_v2.vote.models import (
    AUDIENCE_IDS,
    AudienceAssessmentDraft,
    AudienceReviewPayload,
    ClusterDraft,
    ClusterPayload,
    HotspotRationale,
    HotspotVote,
    VOTER_IDS,
    VoterBallot,
    VoterBallotPayload,
)


class FakeLLM:
    def __init__(self) -> None:
        self.call_count = 0

    def chat_json(self, system: str, user: str, model, **kwargs):
        self.call_count += 1
        name = getattr(model, "__name__", "")
        if name == "ClusterPayload":
            return ClusterPayload(
                clusters=[
                    ClusterDraft(direction="高端影像升级", hotspot_ids=[1, 3, 5, 7, 9, 11]),
                    ClusterDraft(direction="超大容量电池与长续航", hotspot_ids=[2, 4, 6, 8, 10, 12]),
                ]
            )
        if name == "AudienceReviewPayload":
            return AudienceReviewPayload(
                assessments=[
                    AudienceAssessmentDraft(
                        hotspot_index=1,
                        relevance=3,
                        why_it_matters="影像升级影响竞争格局",
                        action_implication="跟进长焦路线",
                        caveat="",
                    ),
                    AudienceAssessmentDraft(
                        hotspot_index=2,
                        relevance=2,
                        why_it_matters="续航成为标配门槛",
                        action_implication="核对电池规格",
                        caveat="投影未给销量",
                    ),
                ]
            )
        voter = next((vid for vid in VOTER_IDS if f"热点投票人 {vid}" in system), "AAA")
        return VoterBallotPayload(
            hotspots=[
                HotspotVote(
                    title=f"{voter}-影像",
                    brand="OPPO",
                    model="Find X9 Ultra",
                    rationale=HotspotRationale(innovation="10倍光变"),
                    axis_tags=[voter],
                ),
                HotspotVote(
                    title=f"{voter}-续航",
                    brand="小米",
                    model="17 Max",
                    rationale=HotspotRationale(user="8000mAh 电池"),
                    axis_tags=[voter],
                ),
            ]
        )


def _resolved_payload() -> dict:
    return {
        "schema_version": 1,
        "run_id": "run-vote-1",
        "category": "phone",
        "input_revision": "verify:run-vote-1",
        "items": [
            {
                "device_id": "phone:oppo:findx9ultra",
                "brand": "OPPO",
                "model": "Find X9 Ultra",
                "field_name": "屏幕显示技术",
                "display_text": "2K 120Hz",
                "policy": "unanimous",
            },
            {
                "device_id": "phone:xiaomi:17max",
                "brand": "小米",
                "model": "17 Max",
                "field_name": "功耗与续航管理",
                "display_text": "8000mAh 金沙江电池",
                "policy": "unanimous",
            },
        ],
    }


def test_vote_stage_writes_votes_and_excel(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "data"))
    from terminal_insight_v2.config import Settings
    import terminal_insight_v2.vote.api as vote_api

    s = Settings(_env_file=None, data_root=str(tmp_path / "data"), openai_api_key="test-key")
    monkeypatch.setattr(vote_api, "settings", s)

    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "resolved.json").write_text(
        json.dumps(_resolved_payload(), ensure_ascii=False),
        encoding="utf-8",
    )
    ctx = StageContext(
        run_id="run-vote-1",
        category="phone",
        schema_fingerprint="fp",
        date_start="2026-07-01",
        date_end="2026-08-11",
        input_revision="verify:run-vote-1",
        run_dir=str(run_dir),
    )
    llm = FakeLLM()
    result = run_vote(ctx, llm=llm)
    assert result.status == "ok"
    assert result.metrics["voters"] == 6
    assert result.metrics["hotspots"] == 12
    assert result.metrics["clusters"] == 2
    assert result.metrics["audiences"] == 4
    assert result.metrics["audiences_reviewed"] == 4
    assert llm.call_count == 11

    data = json.loads((run_dir / "votes.json").read_text(encoding="utf-8"))
    assert data["schema_version"] == 3
    assert [v["voter_id"] for v in data["voters"]] == list(VOTER_IDS)
    assert data["clusters"][0]["voter_ids"]
    assert [r["audience_id"] for r in data["audience_reviews"]] == list(AUDIENCE_IDS)
    assert data["audience_reviews"][0]["assessments"][0]["hotspot_index"] == 1
    assert data["audience_reviews"][0]["assessments"][0]["title"] == "AAA-影像"

    wb = load_workbook(result.artifacts["excel"])
    ws = wb[SHEET_NAME]
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.max_row == 13
    aws = wb[AUDIENCE_SHEET_NAME]
    assert [c.value for c in aws[1]] == AUDIENCE_HEADERS
    assert aws.max_row == 9
    assert aws.cell(2, 1).value == "工信/监管"
    assert aws.cell(2, 3).value == "AAA-影像"
