"""Vote + hypothesis mapping integration (dual-pass consensus)."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from terminal_insight_v2.hypothesis.models import (
    HypothesisMappingDraft,
    HypothesisMappingPayload,
)
from terminal_insight_v2.models.stage import StageContext
from terminal_insight_v2.vote.api import run_vote
from terminal_insight_v2.vote.excel_write import HYPOTHESIS_HEADERS, HYPOTHESIS_SHEET_NAME
from terminal_insight_v2.vote.models import (
    AudienceAssessmentDraft,
    AudienceReviewPayload,
    ClusterDraft,
    ClusterPayload,
    HotspotRationale,
    HotspotVote,
    VOTER_IDS,
    VoterBallotPayload,
)


class FakeLLM:
    """Deterministic LLM. Hypothesis mapping can agree or conflict across passes."""

    def __init__(
        self,
        *,
        map_hypotheses: bool = False,
        conflict: bool = False,
    ) -> None:
        self.call_count = 0
        self.map_hypotheses = map_hypotheses
        self.conflict = conflict
        self._hypo_pass = 0

    def chat_json(self, system: str, user: str, model, **kwargs):
        self.call_count += 1
        name = getattr(model, "__name__", "")
        if name == "ClusterPayload":
            return ClusterPayload(
                clusters=[
                    ClusterDraft(direction="高端影像升级", hotspot_ids=[1, 3, 5, 7, 9, 11]),
                    ClusterDraft(
                        direction="超大容量电池与长续航", hotspot_ids=[2, 4, 6, 8, 10, 12]
                    ),
                ]
            )
        if name == "AudienceReviewPayload":
            return AudienceReviewPayload(
                assessments=[
                    AudienceAssessmentDraft(
                        hotspot_index=1,
                        relevance=3,
                        why_it_matters="影像",
                        action_implication="跟进",
                        caveat="",
                    ),
                    AudienceAssessmentDraft(
                        hotspot_index=2,
                        relevance=2,
                        why_it_matters="续航",
                        action_implication="核对",
                        caveat="",
                    ),
                ]
            )
        if name == "HypothesisMappingPayload":
            if not self.map_hypotheses:
                return HypothesisMappingPayload(mappings=[])
            self._hypo_pass += 1
            stance_h1 = "support"
            if self.conflict and self._hypo_pass >= 2:
                stance_h1 = "weaken"
            return HypothesisMappingPayload(
                mappings=[
                    HypothesisMappingDraft(
                        hypothesis_id="H1",
                        hotspot_index=1,
                        stance=stance_h1,  # type: ignore[arg-type]
                        evidence="屏幕显示技术：2K 120Hz",
                        strength=3,
                        note="影像旗舰",
                    ),
                    HypothesisMappingDraft(
                        hypothesis_id="H4a",
                        hotspot_index=2,
                        stance="support",
                        evidence="功耗与续航管理：8000mAh 金沙江电池",
                        strength=2,
                        note="大电池",
                    ),
                ]
            )
        voter = next((vid for vid in VOTER_IDS if f"热点投票人 {vid}" in system), "AAA")
        return VoterBallotPayload(
            hotspots=[
                HotspotVote(
                    title=f"{voter}-影像",
                    brand="OPPO",
                    model="Find X9 Ultra",
                    rationale=HotspotRationale(innovation="10倍光变"),
                    axis_tags=[voter],
                ),
                HotspotVote(
                    title=f"{voter}-续航",
                    brand="小米",
                    model="17 Max",
                    rationale=HotspotRationale(user="8000mAh 电池"),
                    axis_tags=[voter],
                ),
            ]
        )


def _resolved_payload() -> dict:
    return {
        "schema_version": 1,
        "run_id": "run-hypo-1",
        "category": "phone",
        "input_revision": "verify:run-hypo-1",
        "items": [
            {
                "device_id": "phone:oppo:findx9ultra",
                "brand": "OPPO",
                "model": "Find X9 Ultra",
                "field_name": "屏幕显示技术",
                "display_text": "2K 120Hz",
                "policy": "unanimous",
            },
            {
                "device_id": "phone:xiaomi:17max",
                "brand": "小米",
                "model": "17 Max",
                "field_name": "功耗与续航管理",
                "display_text": "8000mAh 金沙江电池",
                "policy": "unanimous",
            },
        ],
    }


def _patch_settings(monkeypatch, tmp_path: Path):
    from terminal_insight_v2.config import Settings
    import terminal_insight_v2.vote.api as vote_api
    import terminal_insight_v2.hypothesis.loader as hypo_loader

    s = Settings(
        _env_file=None,
        data_root=str(tmp_path / "data"),
        openai_api_key="test-key",
    )
    monkeypatch.setattr(vote_api, "settings", s)
    monkeypatch.setattr(hypo_loader, "settings", s)
    return s


def test_vote_skips_hypothesis_llm_when_no_catalog(tmp_path: Path, monkeypatch):
    _patch_settings(monkeypatch, tmp_path)
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "resolved.json").write_text(
        json.dumps(_resolved_payload(), ensure_ascii=False), encoding="utf-8"
    )
    ctx = StageContext(
        run_id="run-hypo-1",
        category="phone",
        schema_fingerprint="fp",
        input_revision="verify:run-hypo-1",
        run_dir=str(run_dir),
    )
    llm = FakeLLM()
    result = run_vote(ctx, llm=llm)
    assert result.status == "ok"
    assert llm.call_count == 11
    assert result.metrics["hypothesis_mappings"] == 0
    assert result.metrics["hypothesis_disputes"] == 0
    assert result.metrics["hypotheses_loaded"] == 0
    data = json.loads((run_dir / "votes.json").read_text(encoding="utf-8"))
    assert data["schema_version"] == 5
    assert data["hypothesis_reviews"] == []
    assert data["hypothesis_disputes"] == []


def test_vote_maps_hypotheses_consensus_and_excel(tmp_path: Path, monkeypatch):
    _patch_settings(monkeypatch, tmp_path)
    catalog = tmp_path / "hypo.json"
    catalog.write_text(
        json.dumps(
            {
                "schema_version": 2,
                "year": 2026,
                "category": "phone",
                "source": "test",
                "hypotheses": [
                    {"id": "H1", "statement": "L3 下探"},
                    {"id": "H4a", "statement": "容量仍是主叙事"},
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "resolved.json").write_text(
        json.dumps(_resolved_payload(), ensure_ascii=False), encoding="utf-8"
    )
    ctx = StageContext(
        run_id="run-hypo-1",
        category="phone",
        schema_fingerprint="fp",
        input_revision="verify:run-hypo-1",
        run_dir=str(run_dir),
        extras={"hypotheses_path": str(catalog)},
    )
    llm = FakeLLM(map_hypotheses=True)
    result = run_vote(ctx, llm=llm)
    assert result.status == "ok"
    assert llm.call_count == 13  # 6+1+4+2 dual-pass
    assert result.metrics["hypothesis_mappings"] == 2
    assert result.metrics["hypothesis_disputes"] == 0
    data = json.loads((run_dir / "votes.json").read_text(encoding="utf-8"))
    assert data["schema_version"] == 5
    assert data["hypothesis_reviews"][0]["hypothesis_id"] == "H1"
    assert data["hypothesis_reviews"][0]["stance"] == "support"
    assert data["hypothesis_reviews"][0]["evidence"]
    assert data["hypothesis_disputes"] == []
    wb = load_workbook(result.artifacts["excel"])
    hws = wb[HYPOTHESIS_SHEET_NAME]
    assert [c.value for c in hws[1]] == HYPOTHESIS_HEADERS
    assert hws.max_row == 3
    assert hws.cell(2, 1).value == "H1"
    assert hws.cell(2, 8).value  # evidence
    assert hws.cell(2, 9).value == 3  # strength


def test_vote_dual_pass_conflict_goes_to_disputes(tmp_path: Path, monkeypatch):
    _patch_settings(monkeypatch, tmp_path)
    catalog = tmp_path / "hypo.json"
    catalog.write_text(
        json.dumps(
            {
                "schema_version": 2,
                "year": 2026,
                "category": "phone",
                "hypotheses": [
                    {"id": "H1", "statement": "L3 下探"},
                    {"id": "H4a", "statement": "容量"},
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "resolved.json").write_text(
        json.dumps(_resolved_payload(), ensure_ascii=False), encoding="utf-8"
    )
    ctx = StageContext(
        run_id="run-hypo-1",
        category="phone",
        schema_fingerprint="fp",
        input_revision="verify:run-hypo-1",
        run_dir=str(run_dir),
        extras={"hypotheses_path": str(catalog)},
    )
    llm = FakeLLM(map_hypotheses=True, conflict=True)
    result = run_vote(ctx, llm=llm)
    assert result.status == "partial"
    assert result.metrics["hypothesis_disputes"] >= 1
    data = json.loads((run_dir / "votes.json").read_text(encoding="utf-8"))
    assert any(d["hypothesis_id"] == "H1" for d in data["hypothesis_disputes"])
    # H4a agreed → still in reviews
    assert any(r["hypothesis_id"] == "H4a" for r in data["hypothesis_reviews"])
    assert any("dual-pass disputes" in e for e in data["errors"])
