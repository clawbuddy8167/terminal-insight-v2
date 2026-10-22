"""Collected-fields xlsx → resolved.json importer."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from terminal_insight_v2.importers.collected_xlsx import (
    format_cell,
    sheet_to_items,
)
from terminal_insight_v2.vote.project import MAX_DEVICES, project_resolved


def test_format_cell_units():
    assert format_cell(2299, kind="price") == "2299元"
    assert format_cell(9000, kind="mah") == "9000mAh"
    assert format_cell(80, kind="watt") == "80W"
    assert format_cell(35.5, kind="weight_g") == "35.5g"
    assert format_cell("2026-01-08") == "2026-01-08"


def test_sheet_to_items_phone_mapping():
    rows = [
        {
            "产品ID": "p1",
            "品牌": "一加",
            "产品名称": "Turbo 6",
            "数据状态": "verified_core",
            "中国区发布日期": "2026-01-08",
            "中国区起售价": 2299,
            "电池容量": 9000,
            "有线充电": 80,
            "无线充电": None,
            "芯片": "骁龙8s",
            "内存规格": "12GB",
            "存储规格": "256GB",
        }
    ]
    items = sheet_to_items(rows, category="phone")
    by_field = {i["field_name"]: i["display_text"] for i in items}
    assert by_field["机型"] == "Turbo 6"
    assert by_field["起售价"] == "2299元"
    assert by_field["电池容量"] == "9000mAh"
    assert "80W" in by_field["充电功率"]
    assert "骁龙8s" in by_field["核心算力与存储"]
    assert items[0]["device_id"] == "p1"


def test_sheet_to_items_skips_candidates():
    rows = [
        {
            "品牌": "荣耀",
            "产品名称": "X5i",
            "数据状态": "candidate_only",
            "中国区起售价": 249,
        }
    ]
    assert sheet_to_items(rows, category="watch") == []
    kept = sheet_to_items(rows, category="watch", include_candidates=True)
    assert len(kept) >= 1


def test_max_devices_covers_96_phones():
    items = []
    for i in range(96):
        items.append(
            {
                "device_id": f"phone:b:m{i}",
                "brand": "B",
                "model": f"M{i}",
                "field_name": "起售价",
                "display_text": "1999元",
                "policy": "unanimous",
            }
        )
    proj = project_resolved(
        {"category": "phone", "items": items},
        source_path=Path("resolved.json"),
    )
    assert MAX_DEVICES >= 96
    assert len(proj.devices) == 96


def test_read_sheet_rows_and_items(tmp_path: Path):
    from terminal_insight_v2.importers.collected_xlsx import read_sheet_rows

    xlsx = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "手机"
    ws["A1"] = "title"
    ws["A4"] = "产品ID"
    ws["B4"] = "品牌"
    ws["C4"] = "产品名称"
    ws["D4"] = "中国区起售价"
    ws["E4"] = "电池容量"
    ws["F4"] = "数据状态"
    ws["A5"] = "id1"
    ws["B5"] = "华为"
    ws["C5"] = "Pura 90"
    ws["D5"] = 4499
    ws["E5"] = 6100
    ws["F5"] = "verified_core"
    wb.save(xlsx)
    rows = read_sheet_rows(xlsx, "手机")
    items = sheet_to_items(rows, category="phone")
    texts = {i["field_name"]: i["display_text"] for i in items}
    assert texts["起售价"] == "4499元"
    assert texts["电池容量"] == "6100mAh"
    assert texts["机型"] == "Pura 90"

